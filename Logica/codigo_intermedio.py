import re
from openpyxl import Workbook
from openpyxl.styles import Font

try:
    from Logica.analizador import (
        reconocimiento, for_header, func_header, patron_return, declaracion_for, asignacion_for, patron_incremento_corto, patron_condicion, operando, tokenizar_expresion, validar_estructura_expresion, tokenizar_condicion, validar_estructura_condicion,
    )
except ImportError:
    from analizador import (
        reconocimiento, for_header, func_header, patron_return, declaracion_for, asignacion_for, patron_incremento_corto, patron_condicion, operando, tokenizar_expresion, validar_estructura_expresion, tokenizar_condicion, validar_estructura_condicion,
    )


PRECEDENCIA = {"*": 5, "/": 5, "+": 4, "-": 4,
               "<": 3, ">": 3, "<=": 3, ">=": 3, "==": 3, "!=": 3,
               "&&": 2, "||": 1}
#PATRON_TOKEN = re.compile(rf"(?P<OPERANDO>{operando})|(?P<OPERADOR>[+\-*/])|(?P<ESPACIO>\s+)")

class Triplos:
    def __init__(self):
        self.triplos = []
        self.posfijos = []
        self.contador_temporales = 0
        self.libres = []
        self.numericos = set()
    def _nuevo_temporal(self):
        if self.libres:
            self.libres.sort(key=lambda t: int(t[1:]))
            return self.libres.pop(0)
        self.contador_temporales +=1
        return f"T{self.contador_temporales}"

    def _liberar(self, temporal):
        self.numericos.discard(temporal)
        self.libres.append(temporal)
    def _es_temporal(self, texto):
        return re.fullmatch(r"T\d+", texto) is not None
    def _emitir(self, objeto, fuente, operador):
        self.triplos.append((objeto, fuente, operador))
    def siguiente_fila(self):
        return len(self.triplos) + 1
    def parchear(self, indices, destino):
        for i in indices:
            objeto, fuente, _ = self.triplos[i]
            self.triplos[i] = (objeto, fuente, destino)
    def _tokenizador(self, valor):
        tokens = tokenizar_expresion(valor)
        if tokens is None or validar_estructura_expresion(tokens):
            raise ValueError(f"Expresión invalida: {valor}")
        return tokens
    def _a_posfija(self, tokens):
        salida = []
        pila_ops = []
        for tipo, texto in tokens:
            if tipo == "OPERANDO":
                salida.append((tipo, texto))
            elif texto == "(":
                pila_ops.append(texto)
            elif texto == ")":
                while pila_ops and pila_ops[-1] != "(":
                    salida.append(("OPERADOR", pila_ops.pop()))
                pila_ops.pop()
            else:
                while pila_ops and pila_ops[-1] != "(" and PRECEDENCIA[pila_ops[-1]] >= PRECEDENCIA[texto]:
                    salida.append(("OPERADOR", pila_ops.pop()))
                pila_ops.append(texto)
        while pila_ops:
            salida.append(("OPERADOR", pila_ops.pop()))
        return salida
    def _generar_desde_posfija(self, posfija):
           pila = []
           for tipo, texto in posfija:
               if tipo == "OPERANDO":
                   pila.append(texto)
                   continue
               derecho = pila.pop()
               izquierdo = pila.pop()
               izq_temp = self._es_temporal(izquierdo)
               der_temp = self._es_temporal(derecho)
               numerico = texto in ("*", "/") or izquierdo in self.numericos or derecho in self.numericos
    
               if izq_temp:
                   resultado = izquierdo
                   self._emitir(resultado, derecho, texto)
                   if der_temp:
                       self._liberar(derecho)
               elif der_temp and (texto == "*" or (texto == "+" and derecho in self.numericos)):
                   resultado = derecho
                   self._emitir(resultado, izquierdo, texto)
               else:
                   resultado = self._nuevo_temporal()
                   self._emitir(resultado, izquierdo, "=")
                   self._emitir(resultado, derecho, texto)
                   if der_temp:
                       self._liberar(derecho)
               if numerico:
                   self.numericos.add(resultado)
               pila.append(resultado)
           return pila.pop()
    
    def traducir_asignacion(self, renglon, destino, valor):
        tokens = self._tokenizador(valor)
        posfija = self._a_posfija(tokens)
        resultado = self._generar_desde_posfija(posfija)
        if not self._es_temporal(resultado):
            temporal = self._nuevo_temporal()
            self._emitir(temporal, resultado, "=")
            resultado = temporal
        self._emitir(destino, resultado, "=")
        self._liberar(resultado)
        texto_posfijo= " ".join(t for _, t in posfija)
        self.posfijos.append((renglon, f"{destino} = {valor}", f"{destino} {texto_posfijo} ="))


    def traducir_condicion(self, renglon, texto_condicion):
        tokens = tokenizar_condicion(texto_condicion)
        if validar_estructura_condicion(tokens):
            raise ValueError(f"Condicion invalida: {texto_condicion}")
        posfija = self._a_posfija(tokens)

        pila = []
        for tipo, texto in posfija:
            if tipo == "OPERANDO":
                pila.append(texto)
            elif texto in ("&&", "||"):
                der = pila.pop()
                izq = pila.pop()
                if texto == "&&":
                    self.parchear(izq["verdaderos"], der["inicio"])
                    pila.append({"inicio": izq["inicio"],
                                 "verdaderos": der["verdaderos"],
                                 "falsos": izq["falsos"] + der["falsos"]})
                else:
                    self.parchear(izq["falsos"], der["inicio"])
                    pila.append({"inicio": izq["inicio"],
                                 "verdaderos": izq["verdaderos"] + der["verdaderos"],
                                 "falsos": der["falsos"]})
            else:
                der = pila.pop()
                izq = pila.pop()
                inicio = self.siguiente_fila()
                temporal = self._nuevo_temporal()
                self._emitir(temporal, izq, "=")
                self._emitir(temporal, der, texto)
                self._emitir("", "TRUE", None)
                fila_true = len(self.triplos) - 1
                self._emitir("", "FALSE", None)
                fila_false = len(self.triplos) - 1
                self._liberar(temporal)
                pila.append({"inicio": inicio, "verdaderos": [fila_true], "falsos": [fila_false]})
        texto_posfijo = " ".join(t for _, t in posfija)
        self.posfijos.append((renglon, texto_condicion, texto_posfijo))
        return pila.pop()
        
def _abrir_for(gen, renglon, m_for):
    init_texto, cond_texto, incr_texto = m_for.group(1), m_for.group(2), m_for.group(3)

    m_init = re.match(declaracion_for, init_texto)
    if m_init:
        gen.traducir_asignacion(renglon, m_init.group(2), m_init.group(3))

    inicio_cond = gen.siguiente_fila()
    fragmento = gen.traducir_condicion(renglon, cond_texto)
    gen.parchear(fragmento["verdaderos"], gen.siguiente_fila())

    return {"renglon": renglon, "incr_texto": incr_texto,
            "inicio_cond": inicio_cond , "falsos": fragmento["falsos"]}

def _cerrar_for(gen, ctx):
    renglon, incr_texto = ctx["renglon"], ctx["incr_texto"]

    m_corto = re.match(patron_incremento_corto, incr_texto)
    if m_corto:
        nombre, op = m_corto.group(1), m_corto.group(2)
        gen.traducir_asignacion(renglon, nombre, f"{nombre} {'+' if op == '++' else '-'} 1")
    else:
        m_larga = re.match(asignacion_for, incr_texto)
        if m_larga:
            gen.traducir_asignacion(renglon, m_larga.group(1), m_larga.group(2))

    gen._emitir("", ctx["inicio_cond"], "JMP")
    gen.parchear(ctx["falsos"], gen.siguiente_fila())


def _escribir_fila(hoja, fila, valores, negrita = False):
    for col, valor in enumerate(valores, start=1):
        if valor == "":
            valor = None
        celda = hoja.cell(row =fila, column=col, value=valor)
        if isinstance(valor, str):
            celda.data_type ="s"
        if negrita:
            celda.font=Font(bold=True)

def guardar_excel(gen, ruta_excel):
    wb = Workbook()

    hoja = wb.active
    hoja.title= "Tabla de Triplos"
    _escribir_fila(hoja, 1, ["#", "Dato objeto", "Dato fuente", "Operador"], negrita=True)
    for i, (objeto, fuente, operador) in enumerate(gen.triplos, start=1):
        _escribir_fila(hoja, i + 1, [i, objeto, fuente, operador])
        
    for letra, ancho in zip("ABCD", (6,16,16,10)):
        hoja.column_dimensions[letra].width = ancho

    fin = len(gen.triplos) + 1
    if any(fuente in ("TRUE", "FALSE") and operador == fin for _, fuente, operador in gen.triplos):
        _escribir_fila(hoja, fin + 1, [fin, None, None, None])

    # hoja2 =wb.create_sheet("posfija")
    # _escribir_fila(hoja2, 1, ["Renglon", "Instruccion", "Notacion posfija"], negrita=True)
    # for i, (renglon, instruccion, posfija) in enumerate(gen.posfijos, start=2):
    #     _escribir_fila(hoja2, i, [renglon, instruccion, posfija])
    # for letra, ancho in zip("ABCD", (9,32,32)):
    #     hoja2.column_dimensions[letra].width = ancho

    wb.save(ruta_excel)

def generar_codigo_intermedio(codigo, tabla_errores=None, ruta_excel="codigo_intermedio.xlsx", solo_si_sin_errores=True):
    errores = tabla_errores.tabla_err() if tabla_errores else []
    if errores and solo_si_sin_errores:
        return None
    lineas_con_error = {e.linea for e in errores}

    gen= Triplos()
    #Cada bloque abierto es un diccionario: tipo ("for"/"funcion"), si se omite, y el contexto del for
    pila_bloques = []

    for renglon, linea in enumerate(codigo.splitlines(), start=1):
        ls = linea.strip()
        en_omision = any(b["omitido"] for b in pila_bloques)   #las funciones aun no se traducen
        traducir = renglon not in lineas_con_error and not en_omision

        m_for = re.match(for_header, ls)
        if m_for:
            if traducir:
                pila_bloques.append({"tipo": "for", "omitido": False, "ctx": _abrir_for(gen, renglon, m_for)})
            else:
                pila_bloques.append({"tipo": "for", "omitido": True, "ctx": None})
            continue
        if re.match(func_header, ls):
            pila_bloques.append({"tipo": "funcion", "omitido": True, "ctx": None})
            continue
        if ls == "}":
            if pila_bloques:
                bloque = pila_bloques.pop()
                if bloque["tipo"] == "for" and not bloque["omitido"]:
                    _cerrar_for(gen, bloque["ctx"])
            continue

        if re.match(patron_return, ls) or not traducir:
            continue

        resultado = reconocimiento(linea)
        try:
            if resultado[0] == "Declaracion":
                _, _tipo, nombre, valor = resultado
                gen.traducir_asignacion(renglon, nombre, valor)
            elif resultado[0] == "Asignacion":
                _, nombre, valor = resultado
                gen.traducir_asignacion(renglon, nombre, valor)
        except ValueError:
            continue

    guardar_excel(gen, ruta_excel)
    return ruta_excel


                